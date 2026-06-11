import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_manual():
    doc = Document()
    
    # Title
    title = doc.add_heading('CHUKA UNIVERSITY CISCO Network Academy', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_heading('Learning Management System (LMS) Training Manual', 1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('\n')
    
    # Content Source
    content_path = 'Training_Manual_Content.md'
    if not os.path.exists(content_path):
        print("Content file not found!")
        return

    with open(content_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # Handle Headings
        if line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=2)
        elif line.startswith('#### '):
            doc.add_heading(line.replace('#### ', ''), level=3)
        elif line.startswith('## '):
            # Already handled subtitle manually or treat as level 1
            if 'Learning Management System' not in line:
                doc.add_heading(line.replace('## ', ''), level=1)
        elif line.startswith('# '):
             if 'CHUKA UNIVERSITY' not in line:
                doc.add_heading(line.replace('# ', ''), level=0)
        
        # Handle Bullet Points
        elif line.startswith('- '):
            p = doc.add_paragraph(line.replace('- ', ''), style='List Bullet')
        
        # Handle Horizontal Rules or Sections
        elif line.startswith('---'):
            doc.add_page_break()
        
        # Regular text
        else:
            p = doc.add_paragraph(line)

    # Footer
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "2026 © Designed by Cisco Academy - Chuka University"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc_name = 'LMS_Training_Manual.docx'
    doc.save(doc_name)
    print(f"Document saved as {doc_name}")

if __name__ == "__main__":
    generate_manual()
